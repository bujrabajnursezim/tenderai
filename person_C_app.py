# -*- coding: utf-8 -*-
import streamlit as st
import pdfplumber
import docx
import plotly.graph_objects as go
from predictor import predict_single
from datetime import datetime
import re

try:
    from fpdf import FPDF
    FPDF_AVAILABLE = True
except ImportError:
    FPDF_AVAILABLE = False

st.set_page_config(page_title="TenderAI", page_icon="🔍", layout="wide")

st.markdown("""
<style>
@import url('https://fonts.googleapis.com/css2?family=Inter:wght@300;400;500;600;700&display=swap');
html, body, [class*="css"] { font-family: 'Inter', sans-serif; background-color: #ffffff; color: #1a1a2e; }
.stApp { background-color: #ffffff; }
.block-container { padding-top: 2rem !important; max-width: 1100px !important; }
.card { background: #f8f8ff; border: 1px solid #e0e0f0; border-radius: 16px; padding: 1.5rem; margin: 0.8rem 0; }
.badge-high { background:#fff0f0; border:1px solid #ff4444; color:#ff4444; padding:6px 18px; border-radius:50px; font-size:0.85rem; font-weight:600; display:inline-block; }
.badge-medium { background:#fff8e8; border:1px solid #ffaa00; color:#cc8800; padding:6px 18px; border-radius:50px; font-size:0.85rem; font-weight:600; display:inline-block; }
.badge-low { background:#f0fff5; border:1px solid #00cc66; color:#009944; padding:6px 18px; border-radius:50px; font-size:0.85rem; font-weight:600; display:inline-block; }
.exp-high { background:#fff5f5; border-left:3px solid #ff4444; padding:0.8rem 1rem; border-radius:0 10px 10px 0; margin:0.4rem 0; font-size:0.88rem; color:#333; }
.exp-medium { background:#fffbf0; border-left:3px solid #ffaa00; padding:0.8rem 1rem; border-radius:0 10px 10px 0; margin:0.4rem 0; font-size:0.88rem; color:#333; }
.exp-low { background:#f5fff8; border-left:3px solid #00cc66; padding:0.8rem 1rem; border-radius:0 10px 10px 0; margin:0.4rem 0; font-size:0.88rem; color:#333; }
.sec { font-size:0.7rem; color:#aaa; text-transform:uppercase; letter-spacing:3px; margin-bottom:0.6rem; margin-top:1.2rem; }
.hist-high { background:#fff5f5; border:1px solid #ffd0d0; border-radius:10px; padding:0.6rem 1rem; margin:0.3rem 0; }
.hist-medium { background:#fffbf0; border:1px solid #ffe0a0; border-radius:10px; padding:0.6rem 1rem; margin:0.3rem 0; }
.hist-low { background:#f5fff8; border:1px solid #b0f0c8; border-radius:10px; padding:0.6rem 1rem; margin:0.3rem 0; }
.highlight-red { background:#ffe0e0; color:#cc2222; padding:1px 4px; border-radius:4px; }
.highlight-orange { background:#fff0cc; color:#cc7700; padding:1px 4px; border-radius:4px; }
.metric-box { background: #f8f8ff; border: 1px solid #e0e0f0; border-radius: 12px; padding: 1rem; text-align: center; }
.metric-val { font-size: 1.6rem; font-weight: 700; color: #6C63FF; }
.metric-lbl { font-size: 0.75rem; color: #999; margin-top: 4px; }
.stProgress > div > div { background: linear-gradient(90deg, #6C63FF, #a78bfa) !important; border-radius: 50px !important; }
</style>
""", unsafe_allow_html=True)

if "history" not in st.session_state:
    st.session_state.history = []

def extract_text(file):
    if file.name.endswith(".pdf"):
        with pdfplumber.open(file) as pdf:
            return " ".join([p.extract_text() or "" for p in pdf.pages])
    elif file.name.endswith(".docx"):
        doc = docx.Document(file)
        return " ".join([p.text for p in doc.paragraphs])
    return ""

def show_gauge(score, height=280):
    color = "#ff4444" if score >= 70 else "#ffaa00" if score >= 40 else "#00cc66"
    fig = go.Figure(go.Indicator(
        mode="gauge+number",
        value=score,
        number={"font": {"size": 44, "color": color}, "suffix": "/100"},
        gauge={
            "axis": {"range": [0, 100], "tickcolor": "#ccc", "tickfont": {"color": "#999", "size": 11}},
            "bar": {"color": color, "thickness": 0.06},
            "bgcolor": "#f8f8ff",
            "bordercolor": "#e0e0f0",
            "steps": [
                {"range": [0, 40], "color": "#e8fff0"},
                {"range": [40, 70], "color": "#fff8e0"},
                {"range": [70, 100], "color": "#fff0f0"},
            ],
            "threshold": {"line": {"color": color, "width": 5}, "thickness": 0.8, "value": score}
        }
    ))
    fig.update_layout(height=height, margin={"t": 10, "b": 0, "l": 20, "r": 20}, paper_bgcolor="#ffffff", font={"color": "#1a1a2e"})
    st.plotly_chart(fig, use_container_width=True)

def highlight_text(text, suspicious_sentences):
    RESTRICTION_WORDS = ["строго", "исключительно", "авторизованный дилер", "уполномоченный дилер", "эксклюзивный"]
    snippet = text[:3000]
    for sent in suspicious_sentences:
        if len(sent) > 10 and sent in snippet:
            snippet = snippet.replace(sent, f'<span class="highlight-red">{sent}</span>')
    for word in RESTRICTION_WORDS:
        pattern = re.compile(re.escape(word), re.IGNORECASE)
        snippet = pattern.sub(f'<span class="highlight-orange">{word}</span>', snippet)
    return snippet

st.markdown("""
<div style='text-align:center;padding:20px 0 10px 0;'>
  <span style='font-size:42px;font-weight:900;background:linear-gradient(90deg,#6C63FF,#00D4FF);-webkit-background-clip:text;-webkit-text-fill-color:transparent;'>TenderAI</span>
  <p style='color:#A0A8C0;font-size:16px;margin-top:5px;'>AI-оценка прозрачности тендеров для МСБ Казахстана</p>
</div>
""", unsafe_allow_html=True)
st.markdown("---")

tab1, tab2, tab3 = st.tabs(["Анализ тендера", "Сравнить два тендера", "История анализов"])

with tab1:
    uploaded_file = st.file_uploader("Загрузите тендерную документацию (PDF или DOCX)", type=["pdf", "docx"], key="main_upload")
    if uploaded_file:
        with st.spinner("Анализируем документ..."):
            text = extract_text(uploaded_file)
        if len(text) < 50:
            st.error("Не удалось извлечь текст из файла")
        else:
            result = predict_single(text)
            score = result["risk_score"]
            st.session_state.history.insert(0, {
                "name": uploaded_file.name,
                "preview": text[:80] + "...",
                "score": score,
                "level": result.get("level", ""),
                "time": datetime.now().strftime("%H:%M:%S")
            })
            st.session_state.history = st.session_state.history[:5]
            st.markdown("---")
            show_gauge(score)

            reqs = result.get("requirements", {})
            labels = result.get("requirement_labels", {})
            if reqs:
                st.markdown("---")
                st.subheader("📋 Что нужно для участия")
                for key, value in reqs.items():
                    icon, label = labels.get(key, ("•", key))
                    st.markdown(f"""
        <div style='background:#1A1D2E; border-left:4px solid #6C63FF;
        padding:10px; border-radius:5px; margin:4px 0; display:flex; gap:10px'>
        <span style='font-size:20px'>{icon}</span>
        <div>
        <span style='color:#A0A8C0; font-size:12px'>{label}</span><br>
        <span style='color:#FFFFFF; font-weight:bold'>{value}</span>
        </div>
        </div>
        """, unsafe_allow_html=True)

            st.markdown("---")
            st.subheader("⚖️ Правовой анализ")

            legal = result["legal"]
            summary = result["legal_summary"]

            if summary["level"] == "critical":
                st.error(f"🔴 {summary['text']}")
            elif summary["level"] == "warning":
                st.warning(f"🟡 {summary['text']}")
            else:
                st.success(f"✅ {summary['text']}")

            for item in legal:
                if item["status"] == "violation":
                    with st.container():
                        st.markdown(f"""
            <div style='background:#2D1B1B; border-left:4px solid #FF4B4B;
            padding:10px; border-radius:5px; margin:5px 0'>
            <b style='color:#FF4B4B'>⚠️ {item['article']}</b><br>
            <span style='color:#FFFFFF'>{item['message']}</span><br>
            <span style='color:#A0A8C0; font-size:12px'>
            Найдено: {item['found']}</span>
            </div>
            """, unsafe_allow_html=True)
                else:
                    with st.container():
                        st.markdown(f"""
            <div style='background:#1B2D1B; border-left:4px solid #00C853;
            padding:10px; border-radius:5px; margin:5px 0'>
            <b style='color:#00C853'>✅ {item['article']}</b><br>
            <span style='color:#A0A8C0'>{item['message']}</span>
            </div>
            """, unsafe_allow_html=True)

            st.caption("⚠️ Анализ носит информационный характер. Рекомендуем проконсультироваться с юристом.")

            if score >= 70:
                st.markdown('<div style="text-align:center"><span class="badge-high">ВЫСОКИЙ РИСК</span></div>', unsafe_allow_html=True)
            elif score >= 40:
                st.markdown('<div style="text-align:center"><span class="badge-medium">СРЕДНИЙ РИСК</span></div>', unsafe_allow_html=True)
            else:
                st.markdown('<div style="text-align:center"><span class="badge-low">НИЗКИЙ РИСК</span></div>', unsafe_allow_html=True)
            st.markdown(f'<p style="text-align:center;color:#555;margin-top:8px;">{result.get("verdict","")}</p>', unsafe_allow_html=True)
            st.markdown(f'<p style="text-align:center;color:#999;font-size:0.85rem;">{result.get("recommendation","")}</p>', unsafe_allow_html=True)
            if result.get("explanations"):
                st.markdown('<div class="sec">AI ОБЪЯСНЕНИЕ</div>', unsafe_allow_html=True)
                for exp in result["explanations"]:
                    css = {"high": "exp-high", "medium": "exp-medium"}.get(exp.get("level",""), "exp-low")
                    st.markdown(f'<div class="{css}">{exp.get("icon","")} {exp.get("text","")}</div>', unsafe_allow_html=True)
            if result.get("price_analysis", {}).get("has_price"):
                st.markdown('<div class="sec">ЦЕНОВОЙ АНАЛИЗ</div>', unsafe_allow_html=True)
                pa = result["price_analysis"]
                st.markdown(f'<div class="exp-medium">{pa["explanation"]}</div>', unsafe_allow_html=True)
                c1, c2 = st.columns(2)
                with c1:
                    st.markdown(f'<div class="metric-box"><div class="metric-val">{pa["tender_price"]:,.0f}</div><div class="metric-lbl">ЦЕНА В ТЕНДЕРЕ (ТГ)</div></div>', unsafe_allow_html=True)
                with c2:
                    st.markdown(f'<div class="metric-box"><div class="metric-val">{pa["market_price"]:,.0f}</div><div class="metric-lbl">РЫНОЧНАЯ ЦЕНА (ТГ)</div></div>', unsafe_allow_html=True)
            if result.get("components"):
                st.markdown('<div class="sec">КОМПОНЕНТЫ РИСКА</div>', unsafe_allow_html=True)
                for name, val in result["components"].items():
                    st.markdown(f'<div style="display:flex;justify-content:space-between;font-size:0.8rem;color:#888;margin-bottom:3px;"><span>{name}</span><span>{val}%</span></div>', unsafe_allow_html=True)
                    st.progress(val / 100)
            if result.get("suspicious_sentences"):
                st.markdown('<div class="sec">ТЕПЛОВАЯ КАРТА ТЕКСТА</div>', unsafe_allow_html=True)
                highlighted = highlight_text(text, result.get("suspicious_sentences", []))
                st.markdown(f'<div class="card" style="font-size:0.82rem;line-height:1.7;color:#333;max-height:300px;overflow-y:auto;">{highlighted}...</div>', unsafe_allow_html=True)
            if result.get("stats"):
                st.markdown('<div class="sec">СТАТИСТИКА</div>', unsafe_allow_html=True)
                stats = result["stats"]
                c1, c2, c3 = st.columns(3)
                with c1:
                    st.markdown(f'<div class="metric-box"><div class="metric-val">{stats.get("total_chars",0):,}</div><div class="metric-lbl">СИМВОЛОВ</div></div>', unsafe_allow_html=True)
                with c2:
                    st.markdown(f'<div class="metric-box"><div class="metric-val">{stats.get("brand_mentions",0)}</div><div class="metric-lbl">БРЕНДОВ</div></div>', unsafe_allow_html=True)
                with c3:
                    st.markdown(f'<div class="metric-box"><div class="metric-val">{stats.get("precise_numbers",0)}</div><div class="metric-lbl">ТОЧНЫХ ПАРАМ.</div></div>', unsafe_allow_html=True)
            if FPDF_AVAILABLE:
                st.markdown('<div class="sec">ОТЧЁТ</div>', unsafe_allow_html=True)
                try:
                    pdf = FPDF()
                    pdf.add_page()
                    pdf.set_font("Helvetica", "B", 20)
                    pdf.set_text_color(108, 99, 255)
                    pdf.cell(0, 12, "TenderAI - Audit Report", ln=True, align="C")
                    pdf.set_font("Helvetica", "", 10)
                    pdf.set_text_color(150, 150, 150)
                    pdf.cell(0, 6, f"Date: {datetime.now().strftime('%d.%m.%Y %H:%M')}", ln=True, align="C")
                    pdf.ln(6)
                    pdf.set_font("Helvetica", "B", 14)
                    color = (255,68,68) if score>=70 else (255,170,0) if score>=40 else (0,204,102)
                    pdf.set_text_color(*color)
                    pdf.cell(0, 10, f"Risk Score: {score}/100", ln=True, align="C")
                    pdf.ln(4)
                    pdf.set_font("Helvetica", "", 10)
                    pdf.set_text_color(50, 50, 80)
                    for exp in result.get("explanations", []):
                        t = exp.get("text","").encode('latin-1','replace').decode('latin-1')
                        pdf.multi_cell(0, 6, f"- {t}")
                    pdf_bytes = pdf.output(dest='S').encode('latin-1')
                    st.download_button(label="Скачать PDF отчёт", data=pdf_bytes, file_name=f"TenderAI_{datetime.now().strftime('%Y%m%d_%H%M')}.pdf", mime="application/pdf")
                except Exception:
                    pass

with tab2:
    st.markdown('<div class="sec">ЗАГРУЗИТЕ ДВА ТЕНДЕРА ДЛЯ СРАВНЕНИЯ</div>', unsafe_allow_html=True)
    c1, c2 = st.columns(2)
    with c1:
        file1 = st.file_uploader("Тендер 1", type=["pdf", "docx"], key="compare1")
    with c2:
        file2 = st.file_uploader("Тендер 2", type=["pdf", "docx"], key="compare2")
    if file1 and file2:
        with st.spinner("Анализируем оба тендера..."):
            r1 = predict_single(extract_text(file1))
            r2 = predict_single(extract_text(file2))
        st.markdown("---")
        c1, c2 = st.columns(2)
        with c1:
            st.markdown(f'<p style="text-align:center;font-weight:600;color:#6C63FF;">{file1.name}</p>', unsafe_allow_html=True)
            show_gauge(r1["risk_score"], height=240)
            s = r1["risk_score"]
            css = "badge-high" if s >= 70 else "badge-medium" if s >= 40 else "badge-low"
            st.markdown(f'<div style="text-align:center"><span class="{css}">{s}/100</span></div>', unsafe_allow_html=True)
        with c2:
            st.markdown(f'<p style="text-align:center;font-weight:600;color:#6C63FF;">{file2.name}</p>', unsafe_allow_html=True)
            show_gauge(r2["risk_score"], height=240)
            s = r2["risk_score"]
            css = "badge-high" if s >= 70 else "badge-medium" if s >= 40 else "badge-low"
            st.markdown(f'<div style="text-align:center"><span class="{css}">{s}/100</span></div>', unsafe_allow_html=True)
        st.markdown("---")
        st.markdown('<div class="sec">ИТОГ СРАВНЕНИЯ</div>', unsafe_allow_html=True)
        diff = r1["risk_score"] - r2["risk_score"]
        if diff > 0:
            st.markdown(f'<div class="exp-high"><b>{file1.name}</b> более рискованный (Risk Score: {r1["risk_score"]}/100)</div>', unsafe_allow_html=True)
        elif diff < 0:
            st.markdown(f'<div class="exp-high"><b>{file2.name}</b> более рискованный (Risk Score: {r2["risk_score"]}/100)</div>', unsafe_allow_html=True)
        else:
            st.markdown('<div class="exp-medium">Оба тендера имеют одинаковый уровень риска</div>', unsafe_allow_html=True)
        if r1.get("components") and r2.get("components"):
            st.markdown('<div class="sec">СРАВНЕНИЕ КОМПОНЕНТОВ</div>', unsafe_allow_html=True)
            for name in r1["components"]:
                v1 = r1["components"].get(name, 0)
                v2 = r2["components"].get(name, 0)
                ca, cb = st.columns(2)
                with ca:
                    color = "#ff4444" if v1 > v2 else "#6C63FF"
                    st.markdown(f'<div style="font-size:0.78rem;color:#888;">{name}: <b style="color:{color}">{v1}%</b></div>', unsafe_allow_html=True)
                    st.progress(v1 / 100)
                with cb:
                    color = "#ff4444" if v2 > v1 else "#6C63FF"
                    st.markdown(f'<div style="font-size:0.78rem;color:#888;">{name}: <b style="color:{color}">{v2}%</b></div>', unsafe_allow_html=True)
                    st.progress(v2 / 100)

with tab3:
    st.markdown('<div class="sec">ПОСЛЕДНИЕ 5 АНАЛИЗОВ</div>', unsafe_allow_html=True)
    if not st.session_state.history:
        st.markdown('<p style="color:#aaa;text-align:center;margin-top:2rem;">История пуста — загрузите тендер для анализа</p>', unsafe_allow_html=True)
    else:
        for item in st.session_state.history:
            s = item["score"]
            css = "hist-high" if s >= 70 else "hist-medium" if s >= 40 else "hist-low"
            score_color = "#ff4444" if s >= 70 else "#cc8800" if s >= 40 else "#009944"
            st.markdown(f"""
            <div class="{css}">
                <div style="display:flex;justify-content:space-between;align-items:center;">
                    <div>
                        <span style="font-weight:600;font-size:0.9rem;">{item['name']}</span>
                        <div style="color:#888;font-size:0.78rem;margin-top:2px;">{item['preview']}</div>
                    </div>
                    <div style="text-align:right;min-width:100px;">
                        <div style="font-size:1.4rem;font-weight:700;color:{score_color}">{s}</div>
                        <div style="color:#bbb;font-size:0.75rem;">{item['time']}</div>
                    </div>
                </div>
            </div>
            """, unsafe_allow_html=True)
    if st.session_state.history:
        if st.button("Очистить историю"):
            st.session_state.history = []
            st.rerun()

st.markdown('<p style="text-align:center;color:#ddd;font-size:0.7rem;margin-top:3rem;">TENDERAI // HACKATHON 2026 // KAZAKHSTAN</p>', unsafe_allow_html=True)
